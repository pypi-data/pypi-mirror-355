"""Microsoft Office document loaders - PowerPoint, Word, Excel."""

from ...core import Attachment, loader
from ... import matchers


@loader(match=matchers.pptx_match)
def pptx_to_python_pptx(att: Attachment) -> Attachment:
    """Load PowerPoint using python-pptx with automatic input source handling."""
    try:
        from pptx import Presentation
        
        # Use the new input_source property - no more repetitive patterns!
        att._obj = Presentation(att.input_source)
            
    except ImportError:
        raise ImportError("python-pptx is required for PowerPoint loading. Install with: pip install python-pptx")
    return att


@loader(match=matchers.docx_match)
def docx_to_python_docx(att: Attachment) -> Attachment:
    """Load Word document using python-docx with automatic input source handling."""
    try:
        from docx import Document
        
        # Use the new input_source property - no more repetitive patterns!
        att._obj = Document(att.input_source)
            
    except ImportError:
        raise ImportError("python-docx is required for Word document loading. Install with: pip install python-docx")
    return att


@loader(match=matchers.excel_match)
def excel_to_openpyxl(att: Attachment) -> Attachment:
    """Load Excel workbook using openpyxl with automatic input source handling."""
    try:
        from openpyxl import load_workbook
        
        # Use the new input_source property - no more repetitive patterns!
        att._obj = load_workbook(att.input_source, read_only=True)
            
    except ImportError:
        raise ImportError("openpyxl is required for Excel loading. Install with: pip install openpyxl")
    return att 